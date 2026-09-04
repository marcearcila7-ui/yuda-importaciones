"""Estado de cuenta del cliente, en Excel y en PDF.

Reproduce el libro con el que se lleva la contabilidad de cada cliente: una fila
por movimiento con su saldo acumulado, agrupada por pedido, y el total al final.
Hasta ahora ese documento solo existia como archivo de Excel llevado a mano; el
sistema tenia los datos pero no habia forma de sacarlos.

Las columnas son las mismas del libro: ENVIO, GUIA, FECHA, DESCRIPCION, VALOR
MERCANCIA, LOGISTICA 5%, ABONO, SALDO y NOTA.
"""
import base64
import logging
from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


logger = logging.getLogger(__name__)

COLS = [
    "ENVÍO",
    "GUÍA",
    "FECHA",
    "DESCRIPCIÓN",
    "VALOR MERCANCÍA",
    "LOGÍSTICA 5%",
    "ABONO",
    "SALDO",
    "NOTA",
]

EMPRESA = "YUDA IMPORTACIONES"
RAZON = "义乌市与达贸易有限公司 · YIWU YUDA TRADING CO.,LTD"

_LOGO: bytes | None = None


def _logo_bytes() -> bytes | None:
    global _LOGO
    if _LOGO is None:
        try:
            _LOGO = (Path(__file__).resolve().parent.parent / "assets" / "logoyuda.png").read_bytes()
        except Exception:
            logger.warning("No se encontro el logo para el estado de cuenta")
            _LOGO = b""
    return _LOGO or None


def _num(valor) -> float:
    return round(float(valor or 0), 2)


def _origen(mov: dict) -> str:
    """"10,000 USDT × 6.65" cuando el abono entro en otra moneda.

    En el libro esto se escribia a mano en la descripcion. Ahora sale del dato
    guardado, asi que no se puede desincronizar del monto.
    """
    monto = mov.get("monto_origen")
    tasa = mov.get("tasa_cambio")
    if monto is None or tasa is None:
        return ""
    moneda = mov.get("moneda_origen") or ""
    return f"{monto:,.2f} {moneda} × {tasa}".strip()


def _descripcion_completa(mov: dict) -> str:
    """Descripcion del movimiento, con el origen del abono si lo hubo."""
    partes = [p for p in (mov.get("descripcion") or "", _origen(mov)) if p]
    return "  ".join(partes)


def generar_estado_cuenta_excel(cuenta: dict, fecha: datetime) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Estado de cuenta"

    fill_empresa = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    fill_head = PatternFill(start_color="4B52E8", end_color="4B52E8", fill_type="solid")
    fill_pedido = PatternFill(start_color="EEF0FD", end_color="EEF0FD", fill_type="solid")
    fill_tot = PatternFill(start_color="0D0D0D", end_color="0D0D0D", fill_type="solid")
    blanco = Font(color="FFFFFF", bold=True)
    centro = Alignment(horizontal="center", vertical="center")

    ncols = len(COLS)
    ultima = get_column_letter(ncols)

    ws.merge_cells(f"A1:{ultima}1")
    ws.row_dimensions[1].height = 42
    logo = _logo_bytes()
    if logo is not None:
        try:
            img = XLImage(BytesIO(logo))
            escala = 46 / img.height
            img.width = round(img.width * escala)
            img.height = 46
            ws.add_image(img, "A1")
        except Exception:
            ws["A1"] = EMPRESA

    ws.merge_cells(f"A2:{ultima}2")
    ws["A2"] = EMPRESA
    ws["A2"].fill = fill_empresa
    ws["A2"].font = Font(color="FFFFFF", bold=True, size=14)
    ws["A2"].alignment = centro
    ws.merge_cells(f"A3:{ultima}3")
    ws["A3"] = RAZON
    ws["A3"].fill = fill_empresa
    ws["A3"].font = Font(color="FFFFFF")
    ws["A3"].alignment = centro

    ws["A5"] = f"CLIENTE: {cuenta['nombre']}"
    ws["A5"].font = Font(bold=True, size=12)
    ws["A6"] = f"NIT: {cuenta.get('nit') or ''}"
    ws["A7"] = f"EMPRESA: {cuenta.get('empresa') or ''}"
    ws["A8"] = f"FECHA: {fecha.strftime('%Y-%m-%d')}"
    ws["A9"] = f"AÑO: {fecha.year}"

    fila = 11
    for pedido in cuenta["pedidos"]:
        titulo = pedido.get("pedido_numero") or "SIN PEDIDO"
        ws.merge_cells(f"A{fila}:{ultima}{fila}")
        ws.cell(row=fila, column=1, value=f"{titulo}   ({pedido['moneda']})")
        ws.cell(row=fila, column=1).fill = fill_pedido
        ws.cell(row=fila, column=1).font = Font(bold=True, color="4B52E8")
        fila += 1

        for idx, titulo_col in enumerate(COLS, start=1):
            celda = ws.cell(row=fila, column=idx, value=titulo_col)
            celda.fill = fill_head
            celda.font = blanco
            celda.alignment = centro
        fila += 1

        for mov in pedido["movimientos"]:
            valores = [
                mov.get("envio"),
                mov.get("guia"),
                mov.get("fecha").strftime("%Y-%m-%d") if mov.get("fecha") else "",
                _descripcion_completa(mov),
                _num(mov.get("valor_mercancia")) or None,
                _num(mov.get("comision_yuda")) or None,
                _num(mov.get("abono")) or None,
                _num(mov.get("saldo")),
                mov.get("nota"),
            ]
            for idx, val in enumerate(valores, start=1):
                ws.cell(row=fila, column=idx, value=val)
            fila += 1

        ws.cell(row=fila, column=4, value="SUBTOTAL")
        ws.cell(row=fila, column=5, value=_num(pedido["compras_totales"]))
        ws.cell(row=fila, column=6, value=_num(pedido["comision_total"]))
        ws.cell(row=fila, column=7, value=_num(pedido["abonos_totales"]))
        ws.cell(row=fila, column=8, value=_num(pedido["saldo_pendiente"]))
        for col in range(1, ncols + 1):
            ws.cell(row=fila, column=col).font = Font(bold=True)
        fila += 2

    # Total del cliente, una linea por moneda: no se mezclan monedas distintas
    for total in cuenta["totales_por_moneda"]:
        ws.cell(row=fila, column=1, value=f"TOTAL {total['moneda']}")
        ws.cell(row=fila, column=5, value=_num(total["compras_totales"]))
        ws.cell(row=fila, column=6, value=_num(total["comision_total"]))
        ws.cell(row=fila, column=7, value=_num(total["abonos_totales"]))
        ws.cell(row=fila, column=8, value=_num(total["saldo_pendiente"]))
        for col in range(1, ncols + 1):
            celda = ws.cell(row=fila, column=col)
            celda.fill = fill_tot
            celda.font = blanco
        fila += 1

    for idx, ancho in enumerate([12, 14, 12, 42, 17, 14, 14, 14, 26], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = ancho

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generar_estado_cuenta_pdf(cuenta: dict, fecha: datetime) -> bytes:
    # Se importa aca y no arriba para que el Excel no dependa de WeasyPrint: asi
    # el estado de cuenta en Excel se puede generar y probar sin esa libreria.
    from app.services.pdf_service import render_pdf

    logo = _logo_bytes()
    logo_html = (
        f'<div class="logo"><img src="data:image/png;base64,{base64.b64encode(logo).decode()}" /></div>'
        if logo
        else ""
    )

    bloques = []
    for pedido in cuenta["pedidos"]:
        titulo = pedido.get("pedido_numero") or "SIN PEDIDO"
        filas = []
        for mov in pedido["movimientos"]:
            filas.append(
                "<tr>"
                f"<td>{mov.get('envio') or ''}</td>"
                f"<td>{mov.get('guia') or ''}</td>"
                f"<td>{mov['fecha'].strftime('%Y-%m-%d') if mov.get('fecha') else ''}</td>"
                f"<td class='desc'>{_descripcion_completa(mov)}</td>"
                f"<td>{_num(mov.get('valor_mercancia')) or ''}</td>"
                f"<td>{_num(mov.get('comision_yuda')) or ''}</td>"
                f"<td>{_num(mov.get('abono')) or ''}</td>"
                f"<td>{_num(mov.get('saldo'))}</td>"
                f"<td class='desc'>{mov.get('nota') or ''}</td>"
                "</tr>"
            )
        filas.append(
            "<tr class='subtotal'>"
            "<td colspan='4'>SUBTOTAL</td>"
            f"<td>{_num(pedido['compras_totales'])}</td>"
            f"<td>{_num(pedido['comision_total'])}</td>"
            f"<td>{_num(pedido['abonos_totales'])}</td>"
            f"<td>{_num(pedido['saldo_pendiente'])}</td>"
            "<td></td></tr>"
        )
        encabezados = "".join(f"<th>{c}</th>" for c in COLS)
        bloques.append(
            f"<div class='pedido'>{titulo} ({pedido['moneda']})</div>"
            f"<table><thead><tr>{encabezados}</tr></thead><tbody>{''.join(filas)}</tbody></table>"
        )

    totales = "".join(
        "<tr>"
        f"<td>TOTAL {t['moneda']}</td>"
        f"<td>{_num(t['compras_totales'])}</td>"
        f"<td>{_num(t['comision_total'])}</td>"
        f"<td>{_num(t['abonos_totales'])}</td>"
        f"<td>{_num(t['saldo_pendiente'])}</td>"
        "</tr>"
        for t in cuenta["totales_por_moneda"]
    )

    html = f"""<!doctype html>
<html><head><meta charset="utf-8"><style>
  @page {{ size: A4 landscape; margin: 1.2cm; }}
  * {{ font-family: Arial, "Noto Sans CJK SC", sans-serif; }}
  body {{ color: #0D0D0D; font-size: 9px; }}
  .logo {{ text-align: center; padding: 4px 0; }}
  .logo img {{ height: 32px; }}
  .empresa {{ background: #1E3A5F; color: #fff; padding: 10px; text-align: center; }}
  .empresa h1 {{ margin: 0; font-size: 16px; }}
  .empresa p {{ margin: 2px 0 0; font-size: 9px; }}
  .datos {{ margin: 10px 0; font-size: 10px; }}
  .pedido {{ background: #EEF0FD; color: #4B52E8; font-weight: bold; padding: 5px 8px;
             margin-top: 10px; border-radius: 6px; }}
  table {{ width: 100%; border-collapse: collapse; margin-top: 4px; }}
  th {{ background: #4B52E8; color: #fff; padding: 4px; font-size: 8px; }}
  td {{ border: 1px solid #E5E7EB; padding: 3px; text-align: center; }}
  td.desc {{ text-align: left; }}
  tr.subtotal td {{ font-weight: bold; background: #F5F5F0; }}
  table.totales td {{ background: #0D0D0D; color: #fff; font-weight: bold; padding: 5px; }}
  tr {{ page-break-inside: avoid; }}
</style></head><body>
  {logo_html}
  <div class="empresa"><h1>{EMPRESA}</h1><p>{RAZON}</p></div>
  <div class="datos">
    <div><strong>CLIENTE:</strong> {cuenta['nombre']}</div>
    <div><strong>NIT:</strong> {cuenta.get('nit') or ''}</div>
    <div><strong>EMPRESA:</strong> {cuenta.get('empresa') or ''}</div>
    <div><strong>FECHA:</strong> {fecha.strftime('%Y-%m-%d')}</div>
  </div>
  {''.join(bloques)}
  <div class="pedido" style="margin-top:14px">TOTAL DEL CLIENTE</div>
  <table class="totales">
    <thead><tr><th>MONEDA</th><th>VALOR MERCANCÍA</th><th>LOGÍSTICA 5%</th><th>ABONOS</th><th>SALDO</th></tr></thead>
    <tbody>{totales}</tbody>
  </table>
</body></html>"""
    return render_pdf(html)
