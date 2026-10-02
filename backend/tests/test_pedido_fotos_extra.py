"""El pedido que se le manda a la tienda lleva TODAS las fotos del producto,
no solo la principal: las extra (más ángulos, o el detalle de un bolso) son
las mismas que vio el cliente en su cotización, y el proveedor las necesita
para no mandar otra cosa.
"""
from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.services.excel_service import generar_formato_pedido
from app.services.imagen_service import fotos_de
from app.services.pdf_service import html_pedido

# 1x1 PNG, lo mínimo que openpyxl acepta como imagen.
PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4"
    "890000000a49444154789c63000100000500010d0a2db40000000049454e44ae426082"
)


class ItemFalso:
    def __init__(self, **kw):
        self.foto_url = None
        self.foto_final_url = None
        self.fotos_extra = None
        self.fotos_extra_final = None
        self.item_no = "A1"
        self.descripcion_es = "Producto"
        self.descripcion_zh = None
        self.descripcion_en = None
        self.ctns = 1
        self.qty_por_ctn = 10
        self.price_rmb = 2.0
        self.largo_cm = self.ancho_cm = self.alto_cm = 10
        self.gw = 1.0
        self.__dict__.update(kw)


def test_fotos_de_pone_la_principal_primero_y_acota_las_extra():
    item = ItemFalso(
        foto_url="p.jpg",
        foto_final_url="p-recortada.jpg",
        fotos_extra={"extra1": "e1.jpg", "extra2": "e2.jpg"},
        # El recorte a mano de una extra gana sobre su original.
        fotos_extra_final={"extra2": "e2-recortada.jpg"},
    )
    assert fotos_de(item) == ["p-recortada.jpg", "e1.jpg", "e2-recortada.jpg"]

    muchas = ItemFalso(foto_url="p.jpg", fotos_extra={f"e{i}": f"{i}.jpg" for i in range(9)})
    assert len(fotos_de(muchas)) == 5  # la principal + 4 extra como tope


def test_el_excel_del_pedido_lleva_la_foto_principal_y_las_extra():
    item = ItemFalso(foto_url="p.jpg", fotos_extra={"extra1": "e1.jpg", "extra2": "e2.jpg"})
    cache = {"p.jpg": PNG, "e1.jpg": PNG, "e2.jpg": PNG}

    xlsx = generar_formato_pedido("Tienda 1468", [item], date(2026, 10, 2), fotos=cache)
    ws = load_workbook(BytesIO(xlsx)).active
    assert len(ws._images) == 3, "deberían ir las 3 fotos, no solo la principal"


def test_un_producto_con_una_sola_foto_no_cambia():
    item = ItemFalso(foto_url="p.jpg")
    xlsx = generar_formato_pedido("Tienda 1468", [item], date(2026, 10, 2), fotos={"p.jpg": PNG})
    ws = load_workbook(BytesIO(xlsx)).active
    assert len(ws._images) == 1


def test_el_pdf_del_pedido_lleva_la_foto_principal_y_las_extra():
    item = ItemFalso(foto_url="p.jpg", fotos_extra={"extra1": "e1.jpg"})
    datauris = {"p.jpg": "data:image/jpeg;base64,AAA", "e1.jpg": "data:image/jpeg;base64,BBB"}

    html = html_pedido("Tienda 1468", [item], date(2026, 10, 2), fotos=datauris)
    assert html.count("<img") == 2, "el PDF debería mostrar las dos fotos"
    assert 'class="varias"' in html
